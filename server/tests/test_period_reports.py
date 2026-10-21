from datetime import date
from io import BytesIO, StringIO
import csv
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from app.db import connect
from app.reports import period_bounds, public_report, report_data
from app.spreadsheet_export import workbook_bytes
from test_server import database
from test_dashboard import people, client_login, event


def test_period_validation_and_inclusive_leap_year():
    start,end=period_bounds(start_date=date(2024,1,1),end_date=date(2024,12,31))
    assert (end-start).days==366 and end.date()==date(2025,1,1)
    for args in ({}, {'start_date':date(2026,9,1)}, {'start_date':date(2026,9,2),'end_date':date(2026,9,1)},
        {'start_date':date(2024,1,1),'end_date':date(2025,1,1)}, {'month':'2026-09','start_date':date(2026,9,1)}):
        with pytest.raises(HTTPException) as error:period_bounds(**args)
        assert error.value.status_code==422


def test_period_cross_month_midnight_and_merged_timeline(people):
    event(people,'employee','2026-08-30T20:59:00Z','2026-08-30T21:00:00Z',keys=999)
    event(people,'employee','2026-08-30T21:00:00Z','2026-08-30T21:01:00Z',keys=117)
    event(people,'employee','2026-08-30T21:01:00Z','2026-08-30T21:03:00Z','afk')
    event(people,'employee','2026-08-30T21:02:00Z','2026-08-30T21:05:00Z','locked')
    event(people,'employee','2026-09-01T20:59:00Z','2026-09-01T21:00:00Z','locked')
    event(people,'employee','2026-09-01T21:00:00Z','2026-09-01T21:01:00Z',keys=999)
    c=client_login();params={'start_date':'2026-08-31','end_date':'2026-09-01','employee':str(people['employee']),'details':'true'}
    r=c.get('/api/v1/reports/period',params=params);assert r.status_code==200,r.text
    data=r.json()
    assert data['totals']==dict(active=60,afk=300,unknown=0,keys=117)
    assert [d['day'] for d in data['days']]==['2026-08-31','2026-09-01']
    assert data['period']==dict(start_date='2026-08-31',end_date='2026-09-01')
    assert 'locked' not in r.text
    segments=data['employees'][0]['days'][0]['segments']
    assert [s['state'] for s in segments]==['active','afk']
    assert segments[1]['end'].startswith('2026-08-30T21:05:00')
    assert data['devices'][0]['id']==str(people['device'])
    # Cached team report returns the same public totals without mutating stored/raw states.
    params.pop('employee');params.pop('details')
    for _ in range(2):assert c.get('/api/v1/reports/period',params=params).json()['totals']==data['totals']
    start,end=period_bounds(start_date=date(2026,8,31),end_date=date(2026,9,1))
    with connect() as db:
        raw=report_data(db,start,end)
        assert raw['totals']['locked']==180 and raw['totals']['afk']==120


@pytest.mark.parametrize('kind',['employees','daily','departments'])
@pytest.mark.parametrize('format',['csv','xlsx'])
def test_export_columns_and_metrics_match_period(people,kind,format):
    event(people,'employee','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z',keys=117)
    event(people,'employee','2026-09-11T12:01:00Z','2026-09-11T12:03:00Z','afk')
    event(people,'employee','2026-09-11T12:03:00Z','2026-09-11T12:06:00Z','locked')
    event(people,'outsider','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z',keys=999)
    c=client_login('head');params=dict(start_date='2026-09-10',end_date='2026-09-12',kind=kind,format=format)
    r=c.get('/api/v1/reports/export',params=params);assert r.status_code==200,r.text
    assert '2026-09-10_2026-09-12.'+format in r.headers['content-disposition']
    if format=='csv':
        rows=list(csv.reader(StringIO(r.content.decode('utf-8-sig')),delimiter=';'))
    else:
        sheet=load_workbook(BytesIO(r.content)).active
        rows=list(sheet.values)
        assert sheet.freeze_panes=='A2'
        assert sheet.cell(2,len(rows[0])).data_type=='n'
    assert all(len(row)==len(rows[0]) for row in rows)
    header=rows[0]
    assert not any('Блокировка' in x for x in header)
    for row in rows[1:]:
        values=dict(zip(header,row))
        assert int(values['Нажатия клавиш'])==117
        assert float(values['AFK (часов)'])==round(300/3600,4)
        assert float(values['Активность (часов)'])==round(60/3600,4)
        assert 'outsider' not in row
    report=c.get('/api/v1/reports/departments',params={k:v for k,v in params.items() if k in ('start_date','end_date')}).json()
    assert report['totals']==dict(active=60,afk=300,unknown=0,keys=117)
    assert all(d['totals']['afk']==300 for d in report['departments'])


def test_period_scope_and_revocation(people):
    event(people,'employee','2026-09-09T12:00:00Z','2026-09-09T12:01:00Z',keys=999)
    event(people,'employee','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z',keys=117)
    c=client_login('head');params=dict(start_date='2026-09-01',end_date='2026-10-02')
    assert c.get('/api/v1/reports/period',params=params).json()['totals']['keys']==117
    assert c.get('/api/v1/reports/period',params={**params,'details':True}).status_code==422
    assert c.get('/api/v1/reports/period',params={**params,'department':str(people['othergroup'])}).status_code==403
    assert c.get('/api/v1/reports/period',params={**params,'employee':str(people['outsider']),'details':True}).status_code==404
    with connect() as db:db.execute('DELETE FROM department_heads')
    assert c.get('/api/v1/reports/period',params=params).json()['employees']==[]
    for format in ['csv','xlsx']:
        assert c.get('/api/v1/reports/export',params={**params,'format':format,'department':str(people['group'])}).status_code==403


def test_legacy_reports_combine_lock_without_changing_raw_event(people):
    event(people,'employee','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z','locked')
    c=client_login()
    for path in ['/reports/month?month=2026-09',f'/reports/day?day=2026-09-11&employee={people["employee"]}']:
        r=c.get('/api/v1'+path);assert r.status_code==200
        assert r.json()['totals']['afk']==60 and 'locked' not in r.text
    with connect() as db:assert db.execute('SELECT state FROM events').fetchone()['state']=='locked'


def test_excel_text_never_becomes_formula():
    sheet=load_workbook(BytesIO(workbook_bytes(['Имя','Нажатия'],[['=1+1',117],['@name',9],[' +SUM(A1)',0]]))).active
    for row in range(2,5):assert sheet.cell(row,1).data_type=='s'
    assert sheet['A2'].value=='=1+1' and sheet['B2'].value==117


def test_department_history_and_full_history_are_distinct(people):
    # The employee was added on Sep 10, but has earlier activity in the month.
    event(people,'employee','2026-09-09T12:00:00Z','2026-09-09T12:01:00Z',keys=371)
    event(people,'employee','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z',keys=133)
    event(people,'outsider','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z',keys=390)
    group=str(people['group']);employee=str(people['employee'])
    params=dict(start_date='2026-09-01',end_date='2026-09-30')
    head=client_login('head')
    assert head.get('/api/v1/reports/period',params={**params,'department':group}).json()['totals']['keys']==133
    admin=client_login()
    # Fill part of the cache before the month; its roster must not erase early history.
    admin.get('/api/v1/reports/period',params=dict(start_date='2026-09-08',end_date='2026-09-09',department=group))
    for _ in range(2):
        all_people=admin.get('/api/v1/reports/period',params=params).json()
        selected=admin.get('/api/v1/reports/period',params={**params,'department':group}).json()
        a=next(p for p in all_people['employees'] if p['id']==employee)
        b=next(p for p in selected['employees'] if p['id']==employee)
        assert a['totals']==dict(active=120,afk=0,unknown=0,keys=504)
        assert b['totals']==dict(active=60,afk=0,unknown=0,keys=133)
        assert all_people['totals']['keys']==894 and selected['totals']['keys']==133
        assert all(p['id']!=str(people['outsider']) for p in selected['employees'])
    details=admin.get('/api/v1/reports/period',params={**params,'department':group,'employee':employee,'details':True}).json()
    assert details['totals']==b['totals']
    assert len(details['employees'][0]['days'])==1 and details['devices']
    assert admin.get('/api/v1/reports/month',params=dict(month='2026-09',department=group)).json()['totals']==b['totals']
    assert admin.get('/api/v1/reports/departments',params={**params,'department':group}).json()['departments'][0]['totals']==b['totals']
    for format in ('csv','xlsx'):
        for kind in ('employees','daily','departments'):
            response=admin.get('/api/v1/reports/export',params={**params,'department':group,'format':format,'kind':kind})
            assert response.status_code==200,response.text
            rows=list(csv.reader(StringIO(response.content.decode('utf-8-sig')),delimiter=';')) if format=='csv' else list(load_workbook(BytesIO(response.content)).active.values)
            index=rows[0].index('Нажатия клавиш')
            values=[int(row[index]) for row in rows[1:]]
            assert values==[133,133] if kind=='departments' else sum(values)==133
    # Admin cache must not leak pre-membership history to a head.
    assert head.get('/api/v1/reports/period',params={**params,'department':group}).json()['totals']['keys']==133


def test_head_full_history_unions_only_authorized_departments(people):
    event(people,'employee','2026-09-09T12:00:00Z','2026-09-09T12:01:00Z',keys=371)
    event(people,'employee','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z',keys=133)
    with connect() as db:
        db.execute("INSERT INTO department_memberships(group_id,employee_id,valid_from) VALUES(%s,%s,'2026-09-01')",(people['othergroup'],people['employee']))
        db.execute('INSERT INTO department_heads(group_id,employee_id) VALUES(%s,%s)',(people['othergroup'],people['head']))
    c=client_login('head');params=dict(start_date='2026-09-01',end_date='2026-09-30')
    a=c.get('/api/v1/reports/period',params=params).json()
    b=c.get('/api/v1/reports/period',params={**params,'department':str(people['group'])}).json()
    assert a['totals']['keys']==504 and b['totals']['keys']==133
    with connect() as db:
        db.execute('DELETE FROM department_heads WHERE group_id=%s',(people['othergroup'],))
    assert c.get('/api/v1/reports/period',params={**params,'department':str(people['group'])}).json()['totals']['keys']==133


def test_former_member_badge_and_new_events_after_exit(people):
    g=str(people['group']);e=str(people['employee'])
    event(people,'employee','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z',keys=133)
    with connect() as db:
        db.execute("UPDATE department_memberships SET valid_to='2026-09-15T08:53:36Z' WHERE group_id=%s AND employee_id=%s",(g,e))
    event(people,'employee','2026-09-15T09:57:00Z','2026-09-15T09:58:00Z',keys=241)
    c=client_login();params={'start_date':'2026-09-01','end_date':'2026-09-30','department':g}
    for _ in range(2):
        data=c.get('/api/v1/reports/period',params=params).json()
        p=next(p for p in data['employees'] if p['id']==e)
        assert p['totals']['keys']==133
        assert p['department_membership']['current'] is False
        assert p['department_membership']['left_at'].startswith('2026-09-15T08:53:36')
    full=c.get('/api/v1/reports/period',params={k:v for k,v in params.items() if k!='department'}).json()
    assert next(p for p in full['employees'] if p['id']==e)['totals']['keys']==374
    with connect() as db:
        db.execute("INSERT INTO department_memberships(group_id,employee_id,valid_from) VALUES(%s,%s,'2026-09-15T10:00:00Z')",(g,e))
    # Badge must refresh even when day-statistic cache remains warm.
    assert c.get('/api/v1/reports/period',params=params).json()['employees'][0]['department_membership']['current'] is True


def test_head_card_after_cached_overview_and_access_denials(people):
    event(people,'employee','2026-09-11T12:00:00Z','2026-09-11T12:01:00Z',keys=133)
    c=client_login('head');base={'start_date':'2026-09-01','end_date':'2026-09-30'}
    for group in ({},{'department':str(people['group'])}):
        for _ in range(2):
            overview=c.get('/api/v1/reports/period',params={**base,**group}).json()
            assert 'segments' not in overview['employees'][0]['days'][0]
            detail=c.get('/api/v1/reports/period',params={**base,**group,'employee':str(people['employee']),'details':True})
            assert detail.status_code==200,detail.text
            assert detail.json()['employees'][0]['days'][0]['segments']
            assert detail.json()['totals']==overview['totals']
        assert c.get('/api/v1/reports/period',params={**base,**group,'employee':str(people['outsider']),'details':True}).status_code==404
    assert c.get('/api/v1/reports/period',params={**base,'department':str(people['othergroup']),'employee':str(people['employee']),'details':True}).status_code==403
