"""Typed Excel cells; strings can never become formulas."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def workbook_bytes(header,rows):
    workbook=Workbook();sheet=workbook.active;sheet.title='Статистика'
    sheet.append(header)
    for row_index,row in enumerate(rows,2):
        if len(row)!=len(header):raise ValueError('export_column_mismatch')
        sheet.append(row)
        for column_index in range(1,len(header)+1):
            cell=sheet.cell(row_index,column_index)
            if isinstance(cell.value,str):cell.data_type='s'
            elif isinstance(cell.value,float):cell.number_format='0.0000'
            elif isinstance(cell.value,int):cell.number_format='0'
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17263E')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[1].height=32
    for index,name in enumerate(header,1):sheet.column_dimensions[get_column_letter(index)].width=min(38,max(19,len(name)+2))
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    stream=BytesIO();workbook.save(stream);return stream.getvalue()
